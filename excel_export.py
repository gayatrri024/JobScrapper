"""excel_export.py - Professional Excel report generator for DevOps Job Hunter.

Exports deduplicated, verified jobs to Excel (.xlsx) using openpyxl with:
- Dedicated output folder: output/jobs_YYYY-MM-DD_HH-MM.xlsx
- Multi-sheet workbook architecture:
  1. 'Verified Opportunities': Verified/Listing Updated dates, Pune or India Remote, 0-3y exp.
  2. 'Unverified Leads': Feed-indexed or estimated dates, Pune or India Remote, 0-3y exp.
  3. 'Rejected & Disqualified': Disqualified roles with explicit rejection reasons.
  4. 'Pipeline & Source Health': Pipeline funnel statistics and source health audit table.
- Exact columns: #, Company, Job Role, Location, Experience, Date Posted, Verification Status,
  Match Score, Priority, Source, Careers Page, LinkedIn / Company URL, DevOps Pillars / Skills, Match Reasons / Gaps
- Clickable hyperlinks on Careers Page and LinkedIn / Company URL
- Bold, professional styled header rows with freeze panes and auto-filters
- Clean sortable tables with auto-adjusted column widths and text wrapping
- Global deduplication prioritizing direct employer application URLs over aggregator links
"""

import json
import logging
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from database import normalize_url, compute_dedup_hash

logger = logging.getLogger("devops_job_hunter.excel_export")

EXCEL_COLUMNS = [
    "#",
    "Company",
    "Job Role",
    "Location",
    "Experience",
    "Date Posted",
    "Verification Status",
    "Match Score",
    "Priority",
    "Source",
    "Careers Page",
    "LinkedIn / Company URL",
    "DevOps Pillars / Skills",
    "Match Reasons / Gaps"
]

REJECTED_COLUMNS = [
    "#",
    "Company",
    "Job Role",
    "Location",
    "Experience Required",
    "Rejection Reason",
    "Match Score",
    "Source",
    "Careers Page"
]


def is_direct_employer_url(url: str) -> bool:
    """Returns True if the URL is a direct employer application or careers portal link rather than a news aggregator redirect."""
    if not url or url.strip() == "" or url.lower() in ["unknown", "none"]:
        return False
    u = url.strip().lower()
    if u.startswith("synthetic://"):
        return False
    if any(agg in u for agg in ["news.google.com", "google.com/url", "google.com/rss"]):
        return False
    return True


def job_quality_tuple(j: Dict[str, Any]) -> Tuple[int, int, float]:
    """Computes a tuple to compare job record quality when deduplicating."""
    url = (j.get("canonical_application_url") or j.get("application_url") or "").strip()
    is_direct = 1 if is_direct_employer_url(url) else 0
    conf = (j.get("date_confidence") or "Unknown").strip().lower()
    conf_score = 2 if conf in ("verified", "listing updated") else (1 if conf in ("high", "medium") else 0)
    score = float(j.get("match_score", 0.0) or 0.0)
    return (is_direct, conf_score, score)


def merge_job_records(winner: Dict[str, Any], loser: Dict[str, Any]) -> Dict[str, Any]:
    """Merges supplemental metadata from the loser record into the winner record."""
    merged = dict(winner)
    for key in ["location", "work_arrangement", "experience_required", "experience_evidence", "posting_date_evidence", "raw_snippet", "source_url"]:
        w_val = (merged.get(key) or "").strip()
        l_val = (loser.get(key) or "").strip()
        if (not w_val or w_val.lower() == "unknown") and (l_val and l_val.lower() != "unknown"):
            merged[key] = l_val

    w_skills = merged.get("relevant_skills") or []
    l_skills = loser.get("relevant_skills") or []
    if isinstance(w_skills, str):
        try:
            w_skills = json.loads(w_skills)
        except Exception:
            w_skills = []
    if isinstance(l_skills, str):
        try:
            l_skills = json.loads(l_skills)
        except Exception:
            l_skills = []
    if not w_skills and l_skills:
        merged["relevant_skills"] = l_skills

    w_url = (merged.get("canonical_application_url") or merged.get("application_url") or "").strip()
    l_url = (loser.get("canonical_application_url") or loser.get("application_url") or "").strip()
    if not is_direct_employer_url(w_url) and is_direct_employer_url(l_url):
        merged["application_url"] = loser.get("application_url") or l_url
        merged["canonical_application_url"] = loser.get("canonical_application_url") or l_url
        merged["source_name"] = loser.get("source_name") or merged.get("source_name")

    return merged


def load_company_targets_map(config_path: Optional[Union[str, Path]] = None) -> Dict[str, Dict[str, Any]]:
    """Loads company_targets.json into a lookup dictionary keyed by lowercase name & aliases."""
    if config_path is None:
        config_path = Path(__file__).resolve().parent / "company_targets.json"
    else:
        config_path = Path(config_path)

    targets_map: Dict[str, Dict[str, Any]] = {}
    if not config_path.exists():
        return targets_map

    try:
        with open(config_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            companies = data.get("companies", [])
            for c in companies:
                name = (c.get("name") or "").strip().lower()
                if name:
                    targets_map[name] = c
                for alias in c.get("aliases", []):
                    a_low = alias.strip().lower()
                    if a_low:
                        targets_map[a_low] = c
    except Exception as e:
        logger.warning(f"Could not load company_targets.json for Excel enrichment: {e}")

    return targets_map


def deduplicate_jobs(jobs: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], int]:
    """
    Deduplicates a list of job dictionaries across all sources and queries:
    - Normalizes URLs (strips tracking params, query fragments, trailing slashes).
    - Checks both normalized canonical URL and normalized (company, title) dedup hash.
    - Prefers direct employer application URLs over aggregator / news links.
    - Prefers verified posting dates and higher match scores.
    - Merges missing fields into the retained record.
    """
    url_to_idx: Dict[str, int] = {}
    hash_to_idx: Dict[str, int] = {}
    deduped: List[Dict[str, Any]] = []
    dupes_removed = 0

    for j in jobs:
        raw_url = (j.get("canonical_application_url") or j.get("application_url") or "").strip()
        norm_url = normalize_url(raw_url) if raw_url.lower().startswith(("http://", "https://")) else ""

        comp = (j.get("company") or "").strip()
        title = (j.get("job_title") or "").strip()
        d_hash = ""
        if comp and title and comp.lower() not in ["unknown", "none"] and title.lower() not in ["unknown", "none"]:
            d_hash = compute_dedup_hash(comp, title)

        match_idx = None
        if norm_url and norm_url in url_to_idx:
            match_idx = url_to_idx[norm_url]
        elif d_hash and d_hash in hash_to_idx:
            match_idx = hash_to_idx[d_hash]

        if match_idx is not None:
            existing_job = deduped[match_idx]
            if job_quality_tuple(j) > job_quality_tuple(existing_job):
                deduped[match_idx] = merge_job_records(winner=j, loser=existing_job)
            else:
                deduped[match_idx] = merge_job_records(winner=existing_job, loser=j)
            dupes_removed += 1
        else:
            new_idx = len(deduped)
            deduped.append(dict(j))
            if norm_url:
                url_to_idx[norm_url] = new_idx
            if d_hash:
                hash_to_idx[d_hash] = new_idx

    return deduped, dupes_removed


def parse_posting_datetime(job: Dict[str, Any]) -> Optional[datetime]:
    """Extracts timezone-aware datetime from job posting_date if available."""
    pdate = (job.get("posting_date") or "").strip()
    if not pdate or pdate.lower() in ["unknown", "unverified"]:
        return None

    try:
        if "t" in pdate.lower():
            clean = pdate.replace("Z", "+00:00")
            return datetime.fromisoformat(clean)
        elif len(pdate) >= 10:
            return datetime.strptime(pdate[:10], "%Y-%m-%d").replace(tzinfo=timezone.utc)
    except Exception:
        pass
    return None


def sort_jobs_by_date(jobs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Sorts jobs by posting date, newest first, when reliable dates are available."""
    def _sort_key(j: Dict[str, Any]) -> Tuple[int, float]:
        dt = parse_posting_datetime(j)
        if dt is not None:
            return (1, dt.timestamp())
        score = j.get("match_score", 0.0) or 0.0
        return (0, float(score))

    return sorted(jobs, key=_sort_key, reverse=True)


def resolve_company_or_linkedin_url(job: Dict[str, Any], targets_map: Dict[str, Dict[str, Any]]) -> str:
    """Resolves a relevant LinkedIn URL or official company career URL without fabricating."""
    for field in ["linkedin_url", "company_url", "official_company_url", "company_careers_url"]:
        val = (job.get(field) or "").strip()
        if val.lower().startswith(("http://", "https://")):
            return val

    src_url = (job.get("source_url") or "").strip()
    if "linkedin.com" in src_url.lower():
        return src_url

    comp = (job.get("company") or "").strip().lower()
    if comp and comp in targets_map:
        target = targets_map[comp]
        c_url = (target.get("career_url") or target.get("website") or "").strip()
        if c_url.lower().startswith(("http://", "https://")):
            return c_url

    return ""


def _format_skills(skills: Any) -> str:
    """Helper to convert skills list or JSON string into a readable comma-separated string."""
    if isinstance(skills, list):
        return ", ".join(str(s) for s in skills if s)
    if isinstance(skills, str):
        if skills.startswith("[") and skills.endswith("]"):
            try:
                parsed = json.loads(skills)
                if isinstance(parsed, list):
                    return ", ".join(str(s) for s in parsed if s)
            except Exception:
                pass
        return skills
    return ""


class ExcelJobExporter:
    """Handles formatting and exporting deduplicated jobs to a multi-tab Excel (.xlsx) workbook."""

    def __init__(self, output_dir: Optional[Union[str, Path]] = None, company_targets_file: Optional[Union[str, Path]] = None):
        if output_dir is None:
            self.output_dir = Path(__file__).resolve().parent / "output"
        else:
            self.output_dir = Path(output_dir)

        self.company_targets = load_company_targets_map(company_targets_file)

    def generate_filepath(self) -> Path:
        """Generates a timestamped file path: output/jobs_YYYY-MM-DD_HH-MM.xlsx."""
        self.output_dir.mkdir(parents=True, exist_ok=True)
        now = datetime.now()
        base_name = f"jobs_{now.strftime('%Y-%m-%d_%H-%M')}"
        target = self.output_dir / f"{base_name}.xlsx"

        counter = 1
        while target.exists():
            target = self.output_dir / f"{base_name}_{counter}.xlsx"
            counter += 1

        return target

    def _write_job_sheet(
        self,
        ws,
        sheet_title: str,
        jobs: List[Dict[str, Any]],
        header_color: str = "1F4E79"
    ) -> None:
        """Formats and writes a standard opportunity sheet (Verified or Leads)."""
        ws.title = sheet_title
        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        regular_font = Font(name="Calibri", size=10)

        ws.append(EXCEL_COLUMNS)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(EXCEL_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        for idx, j in enumerate(jobs, 1):
            row_idx = idx + 1
            ws.row_dimensions[row_idx].height = 22
            is_even = (idx % 2 == 0)

            company_val = (j.get("company") or "").strip()
            if company_val.lower() in ["unknown", "none"]:
                company_val = ""

            role_val = (j.get("job_title") or "").strip()

            location_val = (j.get("location") or "").strip()
            if location_val.lower() in ["unknown", "none"]:
                location_val = ""

            exp_val = (j.get("experience_required") or "").strip()
            if exp_val.lower() in ["unknown", "none"]:
                exp_val = ""

            dt = parse_posting_datetime(j)
            date_val = ""
            if dt is not None and j.get("date_confidence") != "Unknown":
                date_val = dt.strftime("%Y-%m-%d")
            elif j.get("posting_date") and j.get("posting_date") != "Unknown":
                date_val = str(j.get("posting_date"))[:10]

            verif_val = (j.get("date_confidence") or "").strip()
            if verif_val.lower() in ["unknown", "none", ""]:
                verif_val = "Estimated / Unverified"

            match_score_val = round(float(j.get("match_score", 0.0) or 0.0), 1)
            priority_val = (j.get("priority") or "Low").strip()
            source_val = (j.get("source_name") or "Unknown").strip()

            careers_url = ((j.get("canonical_application_url") or j.get("application_url") or "")).strip()
            company_url = resolve_company_or_linkedin_url(j, self.company_targets)
            skills_val = _format_skills(j.get("relevant_skills"))

            reasons = (j.get("match_reasons") or "").strip()
            gaps = (j.get("potential_gaps") or "").strip()
            if reasons and gaps and gaps != "None detected":
                reasons_val = f"{reasons} | Gaps: {gaps}"
            elif reasons:
                reasons_val = reasons
            elif gaps and gaps != "None detected":
                reasons_val = f"Gaps: {gaps}"
            else:
                reasons_val = ""

            cells = [
                ws.cell(row=row_idx, column=1, value=idx),
                ws.cell(row=row_idx, column=2, value=company_val),
                ws.cell(row=row_idx, column=3, value=role_val),
                ws.cell(row=row_idx, column=4, value=location_val),
                ws.cell(row=row_idx, column=5, value=exp_val),
                ws.cell(row=row_idx, column=6, value=date_val),
                ws.cell(row=row_idx, column=7, value=verif_val),
                ws.cell(row=row_idx, column=8, value=match_score_val),
                ws.cell(row=row_idx, column=9, value=priority_val),
                ws.cell(row=row_idx, column=10, value=source_val),
                ws.cell(row=row_idx, column=11),
                ws.cell(row=row_idx, column=12),
                ws.cell(row=row_idx, column=13, value=skills_val),
                ws.cell(row=row_idx, column=14, value=reasons_val),
            ]

            cells[0].alignment = Alignment(horizontal="center", vertical="center")
            cells[1].alignment = Alignment(horizontal="left", vertical="center")
            cells[2].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cells[3].alignment = Alignment(horizontal="left", vertical="center")
            cells[4].alignment = Alignment(horizontal="center", vertical="center")
            cells[5].alignment = Alignment(horizontal="center", vertical="center")
            cells[6].alignment = Alignment(horizontal="center", vertical="center")
            cells[7].alignment = Alignment(horizontal="center", vertical="center")
            cells[8].alignment = Alignment(horizontal="center", vertical="center")
            cells[9].alignment = Alignment(horizontal="center", vertical="center")

            c_app = cells[10]
            if careers_url.lower().startswith(("http://", "https://")):
                c_app.value = careers_url
                c_app.hyperlink = careers_url
                c_app.font = link_font
            else:
                c_app.value = ""
                c_app.font = regular_font
            c_app.alignment = Alignment(horizontal="left", vertical="center")

            c_comp_url = cells[11]
            if company_url.lower().startswith(("http://", "https://")):
                c_comp_url.value = company_url
                c_comp_url.hyperlink = company_url
                c_comp_url.font = link_font
            else:
                c_comp_url.value = ""
                c_comp_url.font = regular_font
            c_comp_url.alignment = Alignment(horizontal="left", vertical="center")

            cells[12].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cells[13].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            for col_i, c in enumerate(cells, 1):
                if col_i not in [11, 12]:
                    c.font = regular_font
                c.border = thin_border
                if is_even:
                    c.fill = alt_row_fill

        ws.freeze_panes = "A2"
        max_row = max(ws.max_row, 1)
        last_col_letter = get_column_letter(len(EXCEL_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max_row}"

        col_width_defaults = {
            1: 6,    # #
            2: 22,   # Company
            3: 30,   # Job Role
            4: 24,   # Location
            5: 16,   # Experience
            6: 14,   # Date Posted
            7: 18,   # Verification Status
            8: 13,   # Match Score
            9: 12,   # Priority
            10: 16,  # Source
            11: 42,  # Careers Page
            12: 32,  # LinkedIn / Company URL
            13: 30,  # DevOps Pillars / Skills
            14: 35   # Match Reasons / Gaps
        }
        for col_idx, default_w in col_width_defaults.items():
            col_letter = get_column_letter(col_idx)
            max_len = default_w
            for cell in ws[col_letter]:
                if cell.value:
                    s = str(cell.value)
                    s_len = min(len(s), 45) if getattr(cell, 'hyperlink', None) else len(s)
                    if s_len > max_len:
                        max_len = s_len
            ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

    def _write_rejected_sheet(
        self,
        ws,
        rejected_jobs: List[Dict[str, Any]]
    ) -> None:
        """Formats and writes the 'Rejected & Disqualified' sheet with explicit exclusion reasons."""
        ws.title = "Rejected & Disqualified"
        header_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        regular_font = Font(name="Calibri", size=10)

        ws.append(REJECTED_COLUMNS)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(REJECTED_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        for idx, j in enumerate(rejected_jobs, 1):
            row_idx = idx + 1
            ws.row_dimensions[row_idx].height = 22
            is_even = (idx % 2 == 0)

            company_val = (j.get("company") or "").strip()
            if company_val.lower() in ["unknown", "none"]:
                company_val = ""
            role_val = (j.get("job_title") or "").strip()
            location_val = (j.get("location") or "").strip()
            if location_val.lower() in ["unknown", "none"]:
                location_val = ""
            exp_val = (j.get("experience_required") or "").strip()
            reason_val = (j.get("rejection_reason") or j.get("potential_gaps") or "Criteria Mismatch").strip()
            score_val = round(float(j.get("match_score", 0.0) or 0.0), 1)
            source_val = (j.get("source_name") or "Unknown").strip()
            careers_url = ((j.get("canonical_application_url") or j.get("application_url") or "")).strip()

            cells = [
                ws.cell(row=row_idx, column=1, value=idx),
                ws.cell(row=row_idx, column=2, value=company_val),
                ws.cell(row=row_idx, column=3, value=role_val),
                ws.cell(row=row_idx, column=4, value=location_val),
                ws.cell(row=row_idx, column=5, value=exp_val),
                ws.cell(row=row_idx, column=6, value=reason_val),
                ws.cell(row=row_idx, column=7, value=score_val),
                ws.cell(row=row_idx, column=8, value=source_val),
                ws.cell(row=row_idx, column=9),
            ]

            cells[0].alignment = Alignment(horizontal="center", vertical="center")
            cells[1].alignment = Alignment(horizontal="left", vertical="center")
            cells[2].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cells[3].alignment = Alignment(horizontal="left", vertical="center")
            cells[4].alignment = Alignment(horizontal="center", vertical="center")
            cells[5].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cells[6].alignment = Alignment(horizontal="center", vertical="center")
            cells[7].alignment = Alignment(horizontal="center", vertical="center")

            c_url = cells[8]
            if careers_url.lower().startswith(("http://", "https://")):
                c_url.value = careers_url
                c_url.hyperlink = careers_url
                c_url.font = link_font
            else:
                c_url.value = ""
                c_url.font = regular_font
            c_url.alignment = Alignment(horizontal="left", vertical="center")

            for col_i, c in enumerate(cells, 1):
                if col_i != 9:
                    c.font = regular_font
                c.border = thin_border
                if is_even:
                    c.fill = alt_row_fill

        ws.freeze_panes = "A2"
        max_row = max(ws.max_row, 1)
        last_col_letter = get_column_letter(len(REJECTED_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max_row}"

        for col_idx in range(1, len(REJECTED_COLUMNS) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 15
            for cell in ws[col_letter]:
                if cell.value:
                    s = str(cell.value)
                    s_len = min(len(s), 45) if getattr(cell, 'hyperlink', None) else len(s)
                    if s_len > max_len:
                        max_len = s_len
            ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

    def _write_health_sheet(
        self,
        ws,
        source_report: Optional[List[Dict[str, Any]]],
        summary_stats: Optional[Dict[str, Any]]
    ) -> None:
        """Formats and writes pipeline funnel statistics and source health audit metrics."""
        ws.title = "Pipeline & Source Health"
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )
        header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        section_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        row_cursor = 1

        # Section 1: Pipeline Summary
        ws.cell(row=row_cursor, column=1, value="Pipeline & Filtering Funnel Summary").font = section_font
        row_cursor += 1

        stats_headers = ["Funnel Metric", "Listing Count", "Description / Filter Logic"]
        ws.row_dimensions[row_cursor].height = 24
        for c_idx, h_text in enumerate(stats_headers, 1):
            c = ws.cell(row=row_cursor, column=c_idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        row_cursor += 1

        stats = summary_stats or {}
        metrics = [
            ("Total Listings Discovered", stats.get("total_fetched", 0), "Total raw job postings retrieved across all enabled sources"),
            ("Duplicates Removed (Cross-Source)", stats.get("dupes_removed", 0), "Duplicate listings consolidated via normalized canonical URL and (company, title) matching"),
            ("Unique Vacancies Evaluated", stats.get("total_unique", 0), "Distinct job vacancies evaluated through scoring and qualification pipeline"),
            ("Verified Opportunities (Sheet 1)", stats.get("verified_count", 0), "Verified or high-confidence posting date, 0-3y experience, Pune or Remote India, score >= 40.0"),
            ("Potentially Relevant Leads (Sheet 2)", stats.get("leads_count", 0), "Feed-indexed or estimated dates, 0-3y experience, Pune or Remote India, score >= 40.0"),
            ("Disqualified & Excluded Listings (Sheet 3)", stats.get("rejected_count", 0), "Senior/Lead (5+y), foreign-only location, non-DevOps role, or expired vacancies")
        ]

        for m_name, m_val, m_desc in metrics:
            ws.row_dimensions[row_cursor].height = 20
            c1 = ws.cell(row=row_cursor, column=1, value=m_name)
            c2 = ws.cell(row=row_cursor, column=2, value=m_val)
            c3 = ws.cell(row=row_cursor, column=3, value=m_desc)
            c1.font = bold_font
            c2.font = bold_font
            c3.font = regular_font
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c3.alignment = Alignment(horizontal="left", vertical="center")
            for c in [c1, c2, c3]:
                c.border = thin_border
            row_cursor += 1

        row_cursor += 2

        # Section 2: Source Health Audit
        ws.cell(row=row_cursor, column=1, value="Public Job Sources Health Audit").font = section_font
        row_cursor += 1

        source_headers = ["Source Name", "Health Status", "Listings Found", "Duration (s)", "Health Notes / Diagnostic Warnings"]
        ws.row_dimensions[row_cursor].height = 24
        for c_idx, h_text in enumerate(source_headers, 1):
            c = ws.cell(row=row_cursor, column=c_idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        row_cursor += 1

        for row_data in (source_report or []):
            ws.row_dimensions[row_cursor].height = 20
            c1 = ws.cell(row=row_cursor, column=1, value=row_data.get("source", ""))
            c2 = ws.cell(row=row_cursor, column=2, value=row_data.get("status", "HEALTHY"))
            c3 = ws.cell(row=row_cursor, column=3, value=row_data.get("jobs_found", 0))
            c4 = ws.cell(row=row_cursor, column=4, value=round(float(row_data.get("duration_sec", 0.0)), 2))
            err_msg = row_data.get("error") or "All queries executed successfully."
            c5 = ws.cell(row=row_cursor, column=5, value=str(err_msg))

            c1.font = regular_font
            c2.font = bold_font
            c3.font = regular_font
            c4.font = regular_font
            c5.font = regular_font

            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c3.alignment = Alignment(horizontal="center", vertical="center")
            c4.alignment = Alignment(horizontal="center", vertical="center")
            c5.alignment = Alignment(horizontal="left", vertical="center")

            status_val = str(row_data.get("status", "")).upper()
            if "FAIL" in status_val:
                c2.font = Font(name="Calibri", size=10, bold=True, color="DC2626")
            elif "PART" in status_val:
                c2.font = Font(name="Calibri", size=10, bold=True, color="D97706")
            else:
                c2.font = Font(name="Calibri", size=10, bold=True, color="16A34A")

            for c in [c1, c2, c3, c4, c5]:
                c.border = thin_border
            row_cursor += 1

        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 60
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 45

    def export_jobs(
        self,
        jobs: List[Dict[str, Any]],
        output_filepath: Optional[Union[str, Path]] = None,
        deduplicate: bool = True,
        leads: Optional[List[Dict[str, Any]]] = None,
        rejected: Optional[List[Dict[str, Any]]] = None,
        source_report: Optional[List[Dict[str, Any]]] = None,
        summary_stats: Optional[Dict[str, Any]] = None
    ) -> Tuple[Path, int, int]:
        """
        Exports jobs to a styled Excel workbook.
        If leads, rejected, or source_report are provided, creates a comprehensive
        multi-tab workbook. Otherwise exports a clean single-tab workbook.
        Returns: (saved_filepath, jobs_exported_count, duplicates_removed_count)
        """
        if not HAS_OPENPYXL:
            raise RuntimeError("openpyxl is not installed. Please install openpyxl to export to Excel.")

        dupes_removed = 0
        if deduplicate:
            jobs, dupes_removed = deduplicate_jobs(jobs)
            if leads:
                leads, _ = deduplicate_jobs(leads)
            if rejected:
                rejected, _ = deduplicate_jobs(rejected)

        sorted_jobs = sort_jobs_by_date(jobs)

        if output_filepath is not None:
            target_path = Path(output_filepath)
            target_path.parent.mkdir(parents=True, exist_ok=True)
        else:
            target_path = self.generate_filepath()

        wb = openpyxl.Workbook()
        ws_primary = wb.active
        self._write_job_sheet(ws_primary, "Verified Opportunities", sorted_jobs, header_color="1F4E79")

        if leads is not None:
            sorted_leads = sort_jobs_by_date(leads)
            ws_leads = wb.create_sheet(title="Unverified Leads")
            self._write_job_sheet(ws_leads, "Unverified Leads", sorted_leads, header_color="0E7490")

        if rejected is not None:
            ws_rej = wb.create_sheet(title="Rejected & Disqualified")
            self._write_rejected_sheet(ws_rej, rejected)

        if source_report is not None or summary_stats is not None:
            ws_health = wb.create_sheet(title="Pipeline & Source Health")
            self._write_health_sheet(ws_health, source_report, summary_stats)

        wb.save(target_path)
        logger.info(f"Exported {len(sorted_jobs)} jobs to Excel: {target_path}")
        return target_path, len(sorted_jobs), dupes_removed

    def export_categorized_jobs(
        self,
        verified_jobs: List[Dict[str, Any]],
        leads: Optional[List[Dict[str, Any]]] = None,
        rejected: Optional[List[Dict[str, Any]]] = None,
        source_report: Optional[List[Dict[str, Any]]] = None,
        summary_stats: Optional[Dict[str, Any]] = None,
        output_filepath: Optional[Union[str, Path]] = None,
        deduplicate: bool = True
    ) -> Tuple[Path, int, int]:
        """Convenience method to export all 4 sheets into the Excel workbook."""
        return self.export_jobs(
            jobs=verified_jobs,
            output_filepath=output_filepath,
            deduplicate=deduplicate,
            leads=leads or [],
            rejected=rejected or [],
            source_report=source_report or [],
            summary_stats=summary_stats or {}
        )
